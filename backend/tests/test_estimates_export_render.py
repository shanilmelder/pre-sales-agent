"""The Estimate export's content and renderers (Story 8.8), without a database: a read model
in, an `.xlsx` and a `.docx` out, each opened again with the library that wrote it."""

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document
from openpyxl import load_workbook

from app.modules.estimates.adapters.export_docx import render_docx
from app.modules.estimates.adapters.export_xlsx import render_xlsx
from app.modules.estimates.application.export_content import (
    ExportFormat,
    build_export,
    export_time,
    file_name,
    slug,
)
from app.modules.estimates.application.models import EstimateVersion
from app.modules.gaps.application.public import Gap

AT = datetime(2026, 10, 5, 14, 32, 59, tzinfo=UTC)
USER = {"id": "0199b1a0-0000-7000-8000-000000000001", "name": "Owner Person"}


def _totals(effort: float, contingency: float) -> dict[str, Any]:
    return {
        "effort_hours": effort,
        "contingency_hours": contingency,
        "total_hours": effort + contingency,
        "role_hours": {"engineer": effort, "project_manager": 0.0, "qa": 0.0},
    }


def _line(n: int, section: str, title: str, effort: float, contingency: float) -> dict[str, Any]:
    return {
        "id": f"0199b1a0-0000-7000-8000-0000000001{n:02d}",
        "position": n,
        "section": section,
        "title": title,
        "basis": f"Basis of {title}",
        "role_mix": {"engineer": 60, "project_manager": 20, "qa": 20},
        "effort_hours": effort,
        "contingency_hours": contingency,
        "total_hours": effort + contingency,
        "role_hours": {"engineer": effort, "project_manager": 0.0, "qa": 0.0},
        "requirements": [],
        "row_version": 1,
        "edited": False,
        "edited_by_name": None,
        "edited_at": None,
        "edit_reason": None,
        "edit_carried_from_version": None,
    }


def _assumption(
    n: int,
    kind: str,
    gap_title: str,
    *,
    hours: float | None = None,
    line: dict[str, Any] | None = None,
    accepted: bool = True,
    carried: int | None = None,
) -> dict[str, Any]:
    return {
        "id": f"0199b1a0-0000-7000-8000-0000000002{n:02d}",
        "kind": kind,
        "wording": f"Wording {n}",
        "amount_hours": hours,
        "line": None if line is None else {"id": line["id"], "title": line["title"]},
        "origin_kind": "gap",
        "origin": {
            "id": f"0199b1a0-0000-7000-8000-0000000003{n:02d}",
            "title": gap_title,
            "category": "other",
            "impact": "high",
            "why_it_matters": "",
            "status": "converted" if accepted else "open",
        },
        "accepted_by": USER if accepted else None,
        "accepted_at": "2026-10-04T09:00:00Z" if accepted else None,
        "row_version": 1,
        "carried_from_version": carried,
    }


SAP = _line(1, "integration", "SAP order interface", 40.0, 8.0)
OPS = _line(2, "non_functional", "24/7 operations", 16.0, 0.0)


def version(unallocated: float = 12.0) -> EstimateVersion:
    return EstimateVersion.model_validate(
        {
            "id": "0199b1a0-0000-7000-8000-000000000099",
            "version": 2,
            "status": "draft",
            "template_version": "demo-1",
            "roles": ["engineer", "project_manager", "qa"],
            "uncovered_count": 0,
            "dropped_count": 0,
            "row_version": 1,
            "created_at": "2026-10-05T10:00:00Z",
            "sections": [
                {"section": "integration", "lines": [SAP], "subtotal": _totals(40.0, 8.0)},
                {"section": "non_functional", "lines": [OPS], "subtotal": _totals(16.0, 0.0)},
            ],
            "totals": _totals(56.0, 8.0 + unallocated),
            "proposal_status": "succeeded",
            "assumptions": {
                "conditions": [
                    _assumption(1, "condition", "WMS version unknown", carried=1),
                    _assumption(2, "condition", "SSO provider", accepted=False),
                ],
                "contingencies": [
                    _assumption(3, "contingency", "Peak order volume", hours=8.0, line=SAP),
                    _assumption(4, "contingency", "Test data owner", hours=unallocated),
                ],
                "contingency_hours": 8.0 + unallocated,
            },
            "counts": {"total": 4, "accepted": 3, "not_accepted": 1},
            "unconverted_gaps": [],
            "unallocated_contingency_hours": unallocated,
            "uncarried_edit_count": 0,
            "uncarried_edits_from_version": None,
        }
    )


def _gap(title: str, status: str, question: str | None, q_status: str = "drafted") -> Gap:
    return Gap.model_validate(
        {
            "id": "0199b1a0-0000-7000-8000-000000000400",
            "title": title,
            "category": "other",
            "trigger": {"kind": "agent_category", "category": "other"},
            "why_it_matters": "",
            "impact": "medium",
            "impact_basis": "",
            "origin": "detected",
            "status": status,
            "converted_to": None,
            "row_version": 1,
            "created_at": "2026-10-05T10:00:00Z",
            "requirements": [],
            "question": None
            if question is None
            else {
                "id": "0199b1a0-0000-7000-8000-000000000500",
                "text": question,
                "topic": "Interfaces",
                "status": q_status,
                "status_changed_at": "2026-10-05T10:00:00Z",
                "row_version": 1,
                "approved_by": {"id": "0199b1a0-0000-7000-8000-000000000600", "name": "Owner"}
                if q_status == "approved"
                else None,
                "approved_at": "2026-10-05T11:00:00Z" if q_status == "approved" else None,
                "edited_by_human": False,
                "last_changed_by": None,
            },
        }
    )


GAPS = [
    _gap("WMS version unknown", "converted", "Which WMS version?", q_status="approved"),
    _gap("Peak order volume", "open", "What is the peak?"),
    _gap("No question", "open", None),
]

HEADER = (
    "Acme Ltd — Pick & Pack",
    "Estimate v2 · Draft — not submitted",
    "Exported 2026-10-05 14:32 UTC",
)


def _document(unallocated: float = 12.0) -> Any:
    return build_export(
        opportunity_title="Acme Ltd — Pick & Pack",
        version=version(unallocated),
        gaps=GAPS,
        exported_at=AT,
    )


def _sheet_rows(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_file_names_are_ascii_slugs() -> None:
    assert slug("Acme Ltd — Pick & Pack") == "acme-ltd-pick-pack"
    assert slug("Société Générale: WMS") == "societe-generale-wms"
    assert slug("———") == "opportunity"
    assert len(slug("x" * 200)) == 60
    assert file_name("Acme Ltd — Pick & Pack", 2, ExportFormat.XLSX) == (
        "acme-ltd-pick-pack-estimate-v2.xlsx"
    )
    assert file_name("Ünïcode", 3, ExportFormat.DOCX).isascii()
    assert export_time(AT) == "2026-10-05 14:32 UTC"


def test_the_workbook_has_three_sheets_each_with_the_header() -> None:
    workbook = load_workbook(BytesIO(render_xlsx(_document())))

    assert workbook.sheetnames == ["Estimate", "Assumptions", "Clarification Questions"]
    for sheet in workbook.worksheets:
        rows = _sheet_rows(sheet)
        assert tuple(r[0] for r in rows[:3]) == HEADER


def test_the_estimate_sheet_copies_the_read_models_numbers() -> None:
    rows = _sheet_rows(load_workbook(BytesIO(render_xlsx(_document()))).worksheets[0])

    body = rows[4:]
    assert body[0] == (
        "Section",
        "Line",
        "Basis",
        "Engineer (%)",
        "PM (%)",
        "QA (%)",
        "Effort (h)",
        "Contingency (h)",
        "Total (h)",
    )
    assert body[1][0] == "Integration"
    assert body[2] == (
        "Integration",
        "SAP order interface",
        "Basis of SAP order interface",
        60,
        20,
        20,
        40.0,
        8.0,
        48.0,
    )
    assert body[3][0] == "Integration subtotal"
    assert body[3][6:] == (40.0, 8.0, 48.0)
    assert body[4][0] == "Non-functional"
    assert body[5][1] == "24/7 operations"
    assert body[6][6:] == (16.0, 0.0, 16.0)
    assert body[7] == ("Unallocated contingency", None, None, None, None, None, None, 12.0, 12.0)
    assert body[8] == ("Total", None, None, None, None, None, 56.0, 20.0, 76.0)
    assert len(body) == 9


def test_no_unallocated_row_when_there_is_none() -> None:
    rows = _sheet_rows(load_workbook(BytesIO(render_xlsx(_document(0.0)))).worksheets[0])

    assert all(r[0] != "Unallocated contingency" for r in rows)
    assert rows[-1] == ("Total", None, None, None, None, None, 56.0, 8.0, 64.0)


def test_the_register_and_the_questions_sheets() -> None:
    workbook = load_workbook(BytesIO(render_xlsx(_document())))
    register = _sheet_rows(workbook["Assumptions"])[4:]
    questions = _sheet_rows(workbook["Clarification Questions"])[4:]

    assert register == [
        ("Kind", "Wording", "Hours", "Linked line", "Origin Gap", "Accepted", "Carried"),
        ("Conditions", None, None, None, None, None, None),
        (
            "Condition",
            "Wording 1",
            None,
            None,
            "WMS version unknown",
            "Accepted by Owner Person, 4 Oct 2026",
            "Carried from v1",
        ),
        ("Condition", "Wording 2", None, None, "SSO provider", "Not accepted", None),
        ("Contingencies", None, None, None, None, None, None),
        (
            "Contingency",
            "Wording 3",
            8.0,
            "SAP order interface",
            "Peak order volume",
            "Accepted by Owner Person, 4 Oct 2026",
            None,
        ),
        (
            "Contingency",
            "Wording 4",
            12.0,
            "Unallocated contingency",
            "Test data owner",
            "Accepted by Owner Person, 4 Oct 2026",
            None,
        ),
    ]
    assert questions == [
        ("Gap", "Impact", "Topic", "Question", "Question status", "Gap status"),
        (
            "WMS version unknown",
            "Medium",
            "Interfaces",
            "Which WMS version?",
            "Approved",
            "Converted",
        ),
        ("Peak order volume", "Medium", "Interfaces", "What is the peak?", "Draft", "Open"),
    ]


def test_the_document_has_the_header_and_the_same_three_tables() -> None:
    doc = Document(BytesIO(render_docx(_document())))

    header = tuple(p.text for p in doc.sections[0].header.paragraphs)
    assert header == HEADER
    headings = [
        p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"
    ]
    assert headings == ["Estimate", "Assumptions", "Clarification Questions"]
    estimate, register, questions = (
        [tuple(c.text for c in row.cells) for row in table.rows] for table in doc.tables
    )
    assert estimate[2] == (
        "Integration",
        "SAP order interface",
        "Basis of SAP order interface",
        "60",
        "20",
        "20",
        "40.0",
        "8.0",
        "48.0",
    )
    assert estimate[-2] == ("Unallocated contingency", "", "", "", "", "", "", "12.0", "12.0")
    assert estimate[-1] == ("Total", "", "", "", "", "", "56.0", "20.0", "76.0")
    assert len(estimate) == 9  # the same rows as the sheet, column row included
    assert register[2][5:] == ("Accepted by Owner Person, 4 Oct 2026", "Carried from v1")
    assert register[3][5] == "Not accepted"
    assert [r[0] for r in questions] == ["Gap", "WMS version unknown", "Peak order volume"]


def test_empty_tables_say_none() -> None:
    empty = version(0.0).model_copy(
        update={
            "assumptions": version(0.0).assumptions.model_copy(
                update={"conditions": [], "contingencies": []}
            )
        }
    )
    document = build_export(
        opportunity_title="Acme",
        version=empty,
        gaps=[_gap("Old", "open", "Superseded?", q_status="superseded")],
        exported_at=AT,
    )
    workbook = load_workbook(BytesIO(render_xlsx(document)))

    assert [r[0] for r in _sheet_rows(workbook["Assumptions"])[5:]] == [
        "Conditions",
        "None",
        "Contingencies",
        "None",
    ]
    assert [r[0] for r in _sheet_rows(workbook["Clarification Questions"])[5:]] == ["None"]


def _with_title(title: str) -> Any:
    v = version()
    line = v.sections[0].lines[0].model_copy(update={"title": title})
    section = v.sections[0].model_copy(update={"lines": [line]})
    return build_export(
        opportunity_title=title,
        version=v.model_copy(update={"sections": [section, *v.sections[1:]]}),
        gaps=GAPS,
        exported_at=AT,
    )


def test_text_that_looks_like_a_formula_stays_text() -> None:
    title = '=HYPERLINK("http://evil.example","Click")'
    workbook = load_workbook(BytesIO(render_xlsx(_with_title(title))))
    sheet = workbook["Estimate"]

    cells = [c for row in sheet.iter_rows() for c in row if c.value == title]
    assert len(cells) == 2  # the header and the line
    assert all(c.data_type == "s" for c in cells)
    assert all(c.data_type != "f" for row in sheet.iter_rows() for c in row)


def test_control_characters_are_dropped_from_the_workbook() -> None:
    workbook = load_workbook(BytesIO(render_xlsx(_with_title("SAP\x0b order\x00 interface\x0c"))))

    rows = _sheet_rows(workbook["Estimate"])
    assert rows[0][0] == "SAP order interface"
    assert rows[6][1] == "SAP order interface"


def test_control_characters_are_dropped_from_the_document() -> None:
    doc = Document(BytesIO(render_docx(_with_title("SAP\x0b order\x00 interface\x0c"))))

    assert doc.sections[0].header.paragraphs[0].text == "SAP order interface"
    assert doc.tables[0].rows[2].cells[1].text == "SAP order interface"


def test_the_column_row_repeats_on_every_page() -> None:
    doc = Document(BytesIO(render_docx(_document())))

    for table in doc.tables:
        assert 'w:tblHeader w:val="true"' in table.rows[0]._tr.xml
        assert "w:tblHeader" not in table.rows[1]._tr.xml
