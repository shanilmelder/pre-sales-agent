"""Importing Actuals from a file (Story 10.2) against a real, migrated Postgres as psa_app,
with the real parse child: the template (one row per Baseline line, traced), upload checks
(415 for a wrong type or content, 413, 403 for anyone but a delivery manager; nothing
stored), the worker job (a corrupt file is `failed` and its job `dead`; Retry), the preview
(every row of the I/O matrix: create, correction reason, bad line, blank line, missing cause,
repeated file), Confirm (2 valid rows written and 1 excluded, 412 on a stale preview, 409 on a
repeated file without the tick, all-or-nothing on a failure), and that the api never parses.

Reuses the Actuals scene of `test_actuals_api`; import jobs are queued a day ahead so no other
claimer takes them, and `drain` runs only this test's.
"""

import asyncio
import csv
import io
import os
import zipfile
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
import sqlalchemy as sa
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.engine import Engine

from app.modules.actuals.application import imports
from app.modules.actuals.application import record as record_module
from app.modules.actuals.domain.imports import COLUMNS
from app.platform.config import Settings
from app.platform.db import create_engine
from app.platform.jobs import queue as q
from app.platform.jobs.registry import JobContext
from app.platform.jobs.runner import LeaseSettings, run_job
from app.platform.parsing import runner
from app.platform.parsing.runner import ParseTimeoutError
from app.platform.uow import unit_of_work
from tests.auth_tokens import auth_app
from tests.conftest import make_client, run_async
from tests.test_actuals_api import (
    CORRECTED,
    DM,
    RECORDED,
    Scene,
    _actuals,
    _baseline,
    _collaborate,
    _line,
    _problem,
    _save,
    _scene,
    _view,
)
from tests.test_gaps_detection import rows
from tests.test_opportunities import BASE, _user
from tests.test_reviews_requests import ENG, _events, _version

COMPLETED = "actuals.import.completed"
TEMPLATE = "actuals.template.exported"
JOB = imports.IMPORT_FILE
LEASE = LeaseSettings(lease_s=30, heartbeat_s=5)
MINE: list[str] = []
HEADERS = [c.header for c in COLUMNS]


@pytest.fixture
def storage_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    root = tmp_path / "files"
    monkeypatch.setenv("PSA_STORAGE_DIR", str(root))
    monkeypatch.setattr(
        imports, "import_settings", lambda: Settings(storage_dir=root, parse_timeout_s=60)
    )
    return root


@pytest.fixture
def client(db_url: str, storage_dir: Path) -> Iterator[TestClient]:
    with make_client(auth_app(db_url)) as c:
        yield c


@pytest.fixture(autouse=True)
def _hidden_jobs(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    real_enqueue = imports.enqueue

    async def hidden_enqueue(uow: Any, payload: Any, **kwargs: Any) -> UUID:
        kwargs["run_after"] = datetime(2999, 1, 1, tzinfo=UTC)
        return await real_enqueue(uow, payload, **kwargs)

    monkeypatch.setattr(imports, "enqueue", hidden_enqueue)
    MINE.clear()
    yield
    _retire(list(MINE))
    MINE.clear()


def _retire(import_ids: list[str]) -> None:
    url = os.environ.get("PSA_DATABASE_URL")
    if not import_ids or not url:
        return
    engine = sa.create_engine(url)
    try:
        with engine.begin() as conn:
            conn.execute(
                sa.text(
                    "UPDATE platform_jobs SET status = 'dead', last_error = 'test: not run' "
                    "WHERE job_type = :t AND status IN ('queued', 'failed_retrying') "
                    "AND payload->>'import_id' = ANY(:i)"
                ),
                {"t": JOB, "i": import_ids},
            )
    finally:
        engine.dispose()


def drain(db_url: str, max_runs: int = 10) -> list[str]:
    """Run this test's import jobs until none is left; retries run at once."""

    async def claim(engine: Any) -> q.ClaimedJob | None:
        async with unit_of_work(engine) as uow:
            row = (
                await uow.session.execute(
                    sa.text(
                        "UPDATE platform_jobs SET status = 'running', lease_owner = 'test-worker', "
                        "lease_expires_at = now() + interval '30 seconds', "
                        "attempts = attempts + 1, updated_at = now() "
                        "WHERE id = (SELECT id FROM platform_jobs WHERE job_type = :t "
                        "AND status IN ('queued', 'failed_retrying') "
                        "AND payload->>'import_id' = ANY(:i) ORDER BY created_at LIMIT 1 "
                        "FOR UPDATE SKIP LOCKED) "
                        "RETURNING id, job_type, payload, attempts, opportunity_id"
                    ),
                    {"t": JOB, "i": list(MINE)},
                )
            ).one_or_none()
        if row is None:
            return None
        return q.ClaimedJob(
            id=row.id,
            job_type=row.job_type,
            payload=row.payload,
            attempts=row.attempts,
            opportunity_id=row.opportunity_id,
            lease_owner="test-worker",
        )

    async def scenario() -> list[str]:
        engine = create_engine(Settings(database_url=db_url))
        outcomes: list[str] = []
        try:
            for _ in range(max_runs):
                job = await claim(engine)
                if job is None:
                    break
                outcomes.append(await run_job(engine, job, LEASE))
        finally:
            await engine.dispose()
        return outcomes

    return run_async(scenario())


# --- helpers --------------------------------------------------------------------------------


def _three_lines(client: TestClient, engine: Engine) -> tuple[Scene, str]:
    """The Actuals scene (L1 36 h, L2 20 h) plus L3 (12 h), with v3 as the Baseline."""
    s = _scene(client, engine, baseline=False)
    l3 = _line(engine, s.v3, 3, "Label printing", "12.0")
    _baseline(engine, s.opp_id, s.v3, s.owner_id)
    return s, l3


def _csv(rows_: list[dict[str, str]]) -> bytes:
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=HEADERS)
    writer.writeheader()
    for row in rows_:
        writer.writerow(row)
    return out.getvalue().encode("utf-8")


def _upload(
    client: TestClient, headers: dict[str, str], opp_id: str, name: str, data: bytes
) -> Any:
    return client.post(
        f"{BASE}/{opp_id}/actuals/imports",
        headers=headers,
        files={"file": (name, data, "application/octet-stream")},
    )


def _import(client: TestClient, s: Scene, name: str, data: bytes, db_url: str) -> dict[str, Any]:
    """Upload, run the job, and return the import as read now."""
    resp = _upload(client, s.dm, s.opp_id, name, data)
    assert resp.status_code == 202, resp.text
    body: dict[str, Any] = resp.json()
    assert body["status"] == "queued"
    MINE.append(body["id"])
    drain(db_url)
    return _get(client, s, body["id"])


def _get(client: TestClient, s: Scene, import_id: str) -> dict[str, Any]:
    resp = client.get(f"{BASE}/{s.opp_id}/actuals/imports/{import_id}", headers=s.dm)
    assert resp.status_code == 200, resp.text
    return resp.json()  # type: ignore[no-any-return]


def _apply(
    client: TestClient,
    s: Scene,
    preview: dict[str, Any],
    *,
    excluded: list[int] | None = None,
    reason: str | None = None,
    confirm_repeat: bool = False,
    row_version: int | None = None,
) -> Any:
    body: dict[str, Any] = {"excluded_rows": excluded or [], "confirm_repeat": confirm_repeat}
    if reason is not None:
        body["reason"] = reason
    tag = preview["row_version"] if row_version is None else row_version
    return client.post(
        f"{BASE}/{s.opp_id}/actuals/imports/{preview['id']}/apply",
        headers={**s.dm, "If-Match": f'"{tag}"'},
        json=body,
    )


def _row(preview: dict[str, Any], number: int) -> dict[str, Any]:
    return next(r for r in preview["rows"] if r["row"] == number)


def _template_bytes(client: TestClient, s: Scene) -> bytes:
    resp = client.get(f"{BASE}/{s.opp_id}/actuals/template", headers=s.dm)
    assert resp.status_code == 200, resp.text
    link = resp.json()
    assert link["file_name"] == "actuals-template-estimate-v3.xlsx"
    downloaded = client.get(link["url"], headers=s.dm)
    assert downloaded.status_code == 200, downloaded.text
    return downloaded.content


def _imports(engine: Engine, opp_id: str) -> list[dict[str, Any]]:
    return rows(engine, "SELECT * FROM actuals_imports WHERE opportunity_id = :o", o=opp_id)


def _blob_count(storage_dir: Path) -> int:
    root = storage_dir / "sha256"
    return 0 if not root.exists() else sum(1 for p in root.rglob("*") if p.is_file())


# --- the template ---------------------------------------------------------------------------


def test_template_has_one_row_per_line_and_is_traced(
    client: TestClient, sync_engine: Engine
) -> None:
    s, l3 = _three_lines(client, sync_engine)
    data = _template_bytes(client, s)
    sheet = load_workbook(io.BytesIO(data), read_only=True).worksheets[0]
    table = [list(r) for r in sheet.iter_rows(values_only=True)]
    assert table[0][: len(HEADERS)] == HEADERS
    assert [r[:5] for r in table[1:]] == [
        [s.l1, "L1", "ERP integration", None, 36],
        [s.l2, "L2", None, None, 20],
        [l3, "L3", None, None, 12],
    ]
    assert all(cell is None for r in table[1:] for cell in r[5:])
    (event,) = _events(sync_engine, s.opp_id, TEMPLATE)
    assert event["subject_type"] == "opportunities.opportunity"
    assert str(event["subject_id"]) == s.opp_id
    assert event["payload"]["line_count"] == 3


def test_template_needs_a_baseline(client: TestClient, sync_engine: Engine) -> None:
    s = _scene(client, sync_engine, baseline=False)
    resp = client.get(f"{BASE}/{s.opp_id}/actuals/template", headers=s.dm)
    _problem(resp, 409, "actuals_no_baseline")


# --- the I/O matrix -------------------------------------------------------------------------


def test_happy_path_from_the_template(client: TestClient, sync_engine: Engine, db_url: str) -> None:
    s, l3 = _three_lines(client, sync_engine)
    workbook = load_workbook(io.BytesIO(_template_bytes(client, s)))
    sheet = workbook.worksheets[0]
    actual_h = HEADERS.index("Actual h") + 1
    cause = HEADERS.index("Variance cause") + 1
    failed = HEADERS.index("Failed Assumptions") + 1
    sheet.cell(row=2, column=actual_h, value=41)
    sheet.cell(row=2, column=cause, value="Estimation error")
    sheet.cell(row=2, column=failed, value="A1")
    sheet.cell(row=3, column=actual_h, value=20)
    sheet.cell(row=4, column=actual_h, value="10.5")
    sheet.cell(row=4, column=cause, value="customer_delay")
    out = io.BytesIO()
    workbook.save(out)

    preview = _import(client, s, "actuals.xlsx", out.getvalue(), db_url)
    assert preview["status"] == "parsed"
    assert preview["estimate_version"] == 3
    assert preview["counts"] == {"create": 3, "correct": 0, "invalid": 0, "skipped": 0}
    assert _row(preview, 2)["line_label"] == "L1"
    assert _actuals(sync_engine, s.opp_id) == []  # nothing before Confirm

    resp = _apply(client, s, preview)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["imported"]["status"] == "applied"
    assert body["imported"]["applied"] == {
        "created": 3,
        "corrected": 0,
        "excluded": 0,
        "skipped": 0,
    }
    assert resp.headers["ETag"] == f'"{body["actuals"]["row_version"]}"'
    lines = {line["line_id"]: line for line in body["actuals"]["lines"]}
    assert lines[s.l1]["variance"]["label"] == "+5.0 h (13.9%)"
    assert lines[s.l1]["actual"]["failed_assumptions"][0]["assumption_id"] == s.assumption
    assert lines[l3]["actual"]["effort_hours"] == 10.5
    assert len(_events(sync_engine, s.opp_id, RECORDED)) == 3
    (completed,) = _events(sync_engine, s.opp_id, COMPLETED)
    payload = completed["payload"]
    assert completed["subject_type"] == "actuals.import"
    assert payload["file_sha256"] == _imports(sync_engine, s.opp_id)[0]["file_sha256"]
    assert (payload["created_count"], payload["corrected_count"]) == (3, 0)
    assert (payload["excluded_count"], payload["skipped_count"]) == (0, 0)

    again = _apply(client, s, preview, row_version=body["actuals"]["row_version"])
    _problem(again, 409, "actuals_import_applied")


def test_two_valid_rows_and_one_excluded(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    """Acceptance: exclude the invalid row and confirm: exactly 2 Actuals, 1 excluded."""
    s = _scene(client, sync_engine)
    other = _version(sync_engine, s.opp_id, 2, status="approved", created_by=s.owner_id)
    foreign = _line(sync_engine, other, 1, "Old line", "8.0")
    data = _csv(
        [
            {"Line ID": s.l1, "Actual h": "36"},
            {"Line ID": s.l2, "Actual h": "18", "Variance cause": "Estimation error"},
            {"Line ID": foreign, "Actual h": "8"},
            {"Line ID": "", "Line": "L9"},  # nothing to record: skipped
        ]
    )
    preview = _import(client, s, "actuals.csv", data, db_url)
    assert preview["counts"] == {"create": 2, "correct": 0, "invalid": 1, "skipped": 1}
    assert _row(preview, 4)["problems"] == ["Row 4: line ID not in Baseline v3"]

    blocked = _apply(client, s, preview)
    _problem(blocked, 422, "actual_invalid")
    assert [e["message"] for e in blocked.json()["errors"]] == ["Row 4: line ID not in Baseline v3"]
    assert _actuals(sync_engine, s.opp_id) == []

    resp = _apply(client, s, preview, excluded=[4])
    assert resp.status_code == 200, resp.text
    assert len(_actuals(sync_engine, s.opp_id)) == 2
    (completed,) = _events(sync_engine, s.opp_id, COMPLETED)
    assert completed["payload"]["created_count"] == 2
    assert completed["payload"]["excluded_count"] == 1
    assert completed["payload"]["skipped_count"] == 1


def test_a_correction_needs_a_reason(client: TestClient, sync_engine: Engine, db_url: str) -> None:
    s = _scene(client, sync_engine)
    assert _save(client, s.dm, s.opp_id, [{"line_id": s.l1, "effort_hours": 36}]).status_code == 200
    data = _csv(
        [
            {"Line ID": s.l1, "Actual h": "40", "Variance cause": "Scope change"},
            {"Line ID": s.l2, "Actual h": "20"},
        ]
    )
    preview = _import(client, s, "actuals.csv", data, db_url)
    row = _row(preview, 2)
    assert (row["status"], row["reason_missing"]) == ("correct", True)
    assert preview["counts"]["correct"] == 1

    missing = _apply(client, s, preview)
    _problem(missing, 422, "actual_invalid")
    assert [e["message"] for e in missing.json()["errors"]] == ["Row 2: correction reason missing"]
    assert len(_actuals(sync_engine, s.opp_id)) == 1  # nothing written

    resp = _apply(client, s, preview, reason="Timesheets reconciled")
    assert resp.status_code == 200, resp.text
    (corrected,) = _events(sync_engine, s.opp_id, CORRECTED)
    assert corrected["payload"]["reason"] == "Timesheets reconciled"
    assert resp.json()["imported"]["applied"]["corrected"] == 1


def test_a_rows_own_correction_reason_is_enough(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    s = _scene(client, sync_engine)
    assert _save(client, s.dm, s.opp_id, [{"line_id": s.l1, "effort_hours": 36}]).status_code == 200
    data = _csv(
        [
            {
                "Line ID": s.l1,
                "Actual h": "40",
                "Variance cause": "Scope change",
                "Correction reason": "Late timesheets",
            }
        ]
    )
    preview = _import(client, s, "actuals.csv", data, db_url)
    assert _row(preview, 2)["reason_missing"] is False
    resp = _apply(client, s, preview)
    assert resp.status_code == 200, resp.text
    (corrected,) = _events(sync_engine, s.opp_id, CORRECTED)
    assert corrected["payload"]["reason"] == "Late timesheets"


def test_missing_cause_and_blank_line_are_invalid(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    s = _scene(client, sync_engine)
    data = _csv(
        [
            {"Line ID": s.l2, "Actual h": "20"},
            {"Line ID": "", "Actual h": "5", "Variance cause": "Scope change"},
            {"Line ID": s.l1, "Actual h": "41"},
            {"Line ID": s.l2, "Actual h": "lots", "Kind": "Partial", "Failed Assumptions": "A9"},
        ]
    )
    preview = _import(client, s, "actuals.csv", data, db_url)
    assert _row(preview, 3)["problems"] == ["Row 3: line ID missing"]
    assert _row(preview, 4)["problems"] == ["Row 4: variance cause missing for a final Actual"]
    assert _row(preview, 5)["problems"] == [
        "Row 5: kind must be Final or Milestone",
        "Row 5: actual hours must be a number",
        "Row 5: A9 isn't in the Baseline v3 Assumptions Register",
    ]
    assert preview["counts"] == {"create": 1, "correct": 0, "invalid": 3, "skipped": 0}
    resp = _apply(client, s, preview, excluded=[3, 4, 5])
    assert resp.status_code == 200, resp.text
    assert len(_actuals(sync_engine, s.opp_id)) == 1


def test_a_milestone_with_cost_in_major_units(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    s = _scene(client, sync_engine)
    data = _csv(
        [
            {
                "Line ID": s.l1,
                "Kind": "milestone",
                "Milestone label": "Sprint 1",
                "Actual h": "12.5",
                "Duration days": "3",
                "Cost": "1,250.50",
                "Currency": "eur",
            },
            {"Line ID": s.l2, "Actual h": "20", "Cost": "5000", "Currency": "JPY"},
        ]
    )
    preview = _import(client, s, "actuals.csv", data, db_url)
    assert preview["counts"]["create"] == 2, preview
    assert _apply(client, s, preview).status_code == 200
    stored = {str(r["estimate_line_id"]): r for r in _actuals(sync_engine, s.opp_id)}
    assert (stored[s.l1]["kind"], stored[s.l1]["milestone_label"]) == ("milestone", "Sprint 1")
    assert (stored[s.l1]["cost_minor"], stored[s.l1]["currency"]) == (125050, "EUR")
    assert (stored[s.l2]["cost_minor"], stored[s.l2]["currency"]) == (5000, "JPY")


def test_wrong_type_is_415_and_stores_nothing(
    client: TestClient, sync_engine: Engine, storage_dir: Path
) -> None:
    s = _scene(client, sync_engine)
    pdf = b"%PDF-1.7\n1 0 obj << >> endobj\n%%EOF\n"
    renamed = _upload(client, s.dm, s.opp_id, "actuals.xlsx", pdf)
    _problem(renamed, 415, "file_content_mismatch")
    assert renamed.json()["detail"] == "Rejected: the content doesn't match .xlsx"
    _problem(_upload(client, s.dm, s.opp_id, "actuals.pdf", pdf), 415, "file_type_not_allowed")
    binary_csv = _upload(client, s.dm, s.opp_id, "actuals.csv", b"\x00\xff\xfe")
    _problem(binary_csv, 415, "file_content_mismatch")
    assert _imports(sync_engine, s.opp_id) == []
    assert _blob_count(storage_dir) == 0


def test_too_large_is_413(client: TestClient, sync_engine: Engine, storage_dir: Path) -> None:
    s = _scene(client, sync_engine)
    big = b"Line ID,Actual h\n" + b"x" * (5 * 1024 * 1024 + 1)
    resp = _upload(client, s.dm, s.opp_id, "actuals.csv", big)
    _problem(resp, 413, "file_too_large")
    assert _imports(sync_engine, s.opp_id) == []
    assert _blob_count(storage_dir) == 0


def test_a_corrupt_workbook_fails_its_job_and_can_be_retried(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    s = _scene(client, sync_engine)
    corrupt = io.BytesIO()
    with zipfile.ZipFile(corrupt, "w") as archive:
        archive.writestr("xl/workbook.xml", "not a workbook")
    failed = _import(client, s, "actuals.xlsx", corrupt.getvalue(), db_url)
    assert (failed["status"], failed["error_code"]) == ("failed", "unreadable")
    assert failed["error_message"]
    (job,) = rows(
        sync_engine,
        "SELECT status FROM platform_jobs WHERE job_type = :t AND payload->>'import_id' = :i",
        t=JOB,
        i=failed["id"],
    )
    assert job["status"] == "dead"

    # Manual entry still works.
    assert _save(client, s.dm, s.opp_id, [{"line_id": s.l2, "effort_hours": 20}]).status_code == 200

    latest = client.get(f"{BASE}/{s.opp_id}/actuals/imports/latest", headers=s.dm)
    assert latest.json()["latest"]["id"] == failed["id"]

    retried = client.post(f"{BASE}/{s.opp_id}/actuals/imports/{failed['id']}/retry", headers=s.dm)
    assert retried.status_code == 200, retried.text
    assert retried.json()["status"] == "queued"
    drain(db_url)
    assert _get(client, s, failed["id"])["status"] == "failed"

    _problem(_apply(client, s, _get(client, s, failed["id"])), 409, "actuals_import_not_ready")


def test_retry_needs_a_failed_import(client: TestClient, sync_engine: Engine, db_url: str) -> None:
    s = _scene(client, sync_engine)
    preview = _import(client, s, "a.csv", _csv([{"Line ID": s.l1, "Actual h": "36"}]), db_url)
    resp = client.post(f"{BASE}/{s.opp_id}/actuals/imports/{preview['id']}/retry", headers=s.dm)
    _problem(resp, 409, "parse_not_failed")


def test_a_file_without_the_columns_fails(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    s = _scene(client, sync_engine)
    failed = _import(client, s, "a.csv", b"Name,Hours\nL1,36\n", db_url)
    assert (failed["status"], failed["error_code"]) == ("failed", "missing_columns")


def test_a_repeated_file_needs_the_tick(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    s = _scene(client, sync_engine)
    data = _csv([{"Line ID": s.l1, "Actual h": "36", "Correction reason": "Re-import"}])
    first = _import(client, s, "a.csv", data, db_url)
    assert first["repeat_message"] is None
    assert _apply(client, s, first).status_code == 200

    second = _import(client, s, "again.csv", data, db_url)
    applied_at = _imports_by_id(sync_engine, first["id"])["applied_at"].astimezone(UTC)
    assert second["repeat_message"] == (
        f"This file was already imported on {applied_at.day} {applied_at:%b}"
    )
    refused = _apply(client, s, second)
    _problem(refused, 409, "actuals_import_repeated")
    assert len(_events(sync_engine, s.opp_id, COMPLETED)) == 1

    resp = _apply(client, s, second, confirm_repeat=True)
    assert resp.status_code == 200, resp.text
    payloads = [e["payload"] for e in _events(sync_engine, s.opp_id, COMPLETED)]
    assert [p["repeated_file"] for p in payloads] == [False, True]


def _imports_by_id(engine: Engine, import_id: str) -> dict[str, Any]:
    (row,) = rows(engine, "SELECT * FROM actuals_imports WHERE id = :i", i=import_id)
    return row


def test_a_stale_preview_is_412(client: TestClient, sync_engine: Engine, db_url: str) -> None:
    s = _scene(client, sync_engine)
    preview = _import(client, s, "a.csv", _csv([{"Line ID": s.l1, "Actual h": "36"}]), db_url)
    assert _save(client, s.dm, s.opp_id, [{"line_id": s.l2, "effort_hours": 20}]).status_code == 200
    _problem(_apply(client, s, preview), 412, "row_version_mismatch")
    assert len(_actuals(sync_engine, s.opp_id)) == 1
    refreshed = _get(client, s, preview["id"])
    assert refreshed["row_version"] == _view(client, s.dm, s.opp_id)["row_version"]
    assert _apply(client, s, refreshed).status_code == 200


def test_all_or_nothing(
    client: TestClient, sync_engine: Engine, db_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    s = _scene(client, sync_engine)
    data = _csv([{"Line ID": s.l1, "Actual h": "36"}, {"Line ID": s.l2, "Actual h": "20"}])
    preview = _import(client, s, "a.csv", data, db_url)
    real_write = record_module._write
    calls = {"n": 0}

    async def failing_write(*args: Any, **kwargs: Any) -> Any:
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("disk on fire")
        return await real_write(*args, **kwargs)

    monkeypatch.setattr(record_module, "_write", failing_write)
    resp = _apply(client, s, preview)
    assert resp.status_code == 500
    assert calls["n"] == 2
    assert _actuals(sync_engine, s.opp_id) == []
    assert _events(sync_engine, s.opp_id, RECORDED) == []
    assert _events(sync_engine, s.opp_id, COMPLETED) == []
    assert _get(client, s, preview["id"])["status"] == "parsed"


def test_only_delivery_managers_import(
    client: TestClient, sync_engine: Engine, storage_dir: Path
) -> None:
    s = _scene(client, sync_engine)
    reviewer, reviewer_id = _user(client, sync_engine, "Rita Reviewer", ENG)
    _collaborate(client, s.owner, s.opp_id, reviewer_id)
    data = _csv([{"Line ID": s.l1, "Actual h": "36"}])
    _problem(_upload(client, reviewer, s.opp_id, "a.csv", data), 403, "forbidden")
    _problem(_upload(client, s.owner, s.opp_id, "a.csv", data), 403, "forbidden")
    for path in ("template", "imports/latest", f"imports/{uuid4()}"):
        resp = client.get(f"{BASE}/{s.opp_id}/actuals/{path}", headers=reviewer)
        _problem(resp, 403, "forbidden")
    stranger, _ = _user(client, sync_engine, "Stranger DM", DM)
    _problem(_upload(client, stranger, s.opp_id, "a.csv", data), 404, "not_found")
    assert _imports(sync_engine, s.opp_id) == []
    assert _blob_count(storage_dir) == 0


def test_the_api_never_parses(
    client: TestClient, sync_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Upload and preview run without the parse child or a workbook reader: the file stays
    `queued` until the worker reads it."""

    def forbidden(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("the api parsed a file")

    monkeypatch.setattr(runner, "run_parse", forbidden)
    monkeypatch.setattr("openpyxl.load_workbook", forbidden)
    monkeypatch.setattr("openpyxl.reader.excel.load_workbook", forbidden)
    s = _scene(client, sync_engine)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(HEADERS)
    sheet.append([s.l1, "L1", None, None, 36, None, None, 36])
    out = io.BytesIO()
    workbook.save(out)
    resp = _upload(client, s.dm, s.opp_id, "a.xlsx", out.getvalue())
    assert resp.status_code == 202, resp.text
    MINE.append(resp.json()["id"])
    queued = _get(client, s, resp.json()["id"])
    assert (queued["status"], queued["rows"], queued["counts"]) == ("queued", [], None)
    (stored,) = _imports(sync_engine, s.opp_id)
    assert stored["rows"] is None


def test_imports_are_never_deleted(sync_engine: Engine) -> None:
    with sync_engine.connect() as conn:
        allowed = conn.execute(
            sa.text("SELECT has_table_privilege('psa_app', 'actuals_imports', 'DELETE')")
        ).scalar_one()
    assert allowed is False


def test_worker_loads_the_job_type() -> None:
    from app import main_worker

    main_worker.load_job_types()
    assert JOB in main_worker.REGISTRY
    assert asyncio.iscoroutinefunction(imports.import_file)


def test_a_too_long_own_reason_is_invalid_even_with_a_batch_reason(
    client: TestClient, sync_engine: Engine, db_url: str
) -> None:
    s = _scene(client, sync_engine)
    assert _save(client, s.dm, s.opp_id, [{"line_id": s.l1, "effort_hours": 36}]).status_code == 200
    data = _csv(
        [
            {
                "Line ID": s.l1,
                "Actual h": "40",
                "Variance cause": "Scope change",
                "Correction reason": "x" * 1001,
            }
        ]
    )
    preview = _import(client, s, "actuals.csv", data, db_url)
    row = _row(preview, 2)
    assert (row["status"], row["reason_missing"]) == ("invalid", False)
    assert row["problems"] == ["Row 2: correction reason is longer than 1,000 characters"]

    resp = _apply(client, s, preview, reason="Timesheets reconciled")
    _problem(resp, 422, "actual_invalid")
    assert [e["message"] for e in resp.json()["errors"]] == [
        "Row 2: correction reason is longer than 1,000 characters"
    ]
    assert len(_actuals(sync_engine, s.opp_id)) == 1  # only the saved one
    assert _events(sync_engine, s.opp_id, CORRECTED) == []


def test_a_timeout_on_every_attempt_fails_the_import(
    client: TestClient, sync_engine: Engine, db_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    s = _scene(client, sync_engine)

    async def timed_out(*args: Any, **kwargs: Any) -> Any:
        raise ParseTimeoutError(60.0)

    monkeypatch.setattr(runner, "run_parse", timed_out)
    resp = _upload(client, s.dm, s.opp_id, "a.csv", _csv([{"Line ID": s.l1, "Actual h": "36"}]))
    assert resp.status_code == 202, resp.text
    MINE.append(resp.json()["id"])
    assert drain(db_url) == ["failed_retrying", "dead"]
    failed = _get(client, s, resp.json()["id"])
    assert (failed["status"], failed["error_code"]) == ("failed", "timeout")
    (job,) = rows(
        sync_engine,
        "SELECT status, attempts FROM platform_jobs WHERE job_type = :t "
        "AND payload->>'import_id' = :i",
        t=JOB,
        i=failed["id"],
    )
    assert (job["status"], job["attempts"]) == ("dead", 2)


def test_a_cancelled_final_attempt_fails_the_import(
    client: TestClient, sync_engine: Engine, db_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The runner enforces the job timeout by cancelling the handler: the final attempt still
    leaves the import `failed` (`timeout`), not `running`."""
    s = _scene(client, sync_engine)
    resp = _upload(client, s.dm, s.opp_id, "a.csv", _csv([{"Line ID": s.l1, "Actual h": "36"}]))
    assert resp.status_code == 202, resp.text
    import_id = UUID(resp.json()["id"])
    MINE.append(str(import_id))

    async def hang(*args: Any, **kwargs: Any) -> None:
        await asyncio.sleep(60)

    monkeypatch.setattr(imports, "_read", hang)

    async def scenario() -> None:
        engine = create_engine(Settings(database_url=db_url))
        ctx = JobContext(
            job_id=uuid4(),
            job_type=JOB,
            attempt=2,
            max_attempts=2,
            opportunity_id=UUID(s.opp_id),
            engine=engine,
        )
        try:
            with pytest.raises(TimeoutError):
                async with asyncio.timeout(0.5):
                    await imports.import_file(ctx, imports.ImportActualsFile(import_id=import_id))
        finally:
            await engine.dispose()

    run_async(scenario())
    failed = _get(client, s, str(import_id))
    assert (failed["status"], failed["error_code"]) == ("failed", "timeout")
