"""Exporting the Estimate (Story 8.8) against a real, migrated Postgres as psa_app:
`GET …/estimate/export?format=xlsx|docx` builds the file from the same read models as the
Estimate and Gaps tabs. Each file is opened again with the library that wrote it and checked
against `GET …/estimate` and `GET …/gaps`; plus who may, no version, a bad format, the trace
event and log privacy.

Reuses the fixtures of `test_intake_extraction.py` and the Assumption helpers of
`test_estimates_assumptions.py`.
"""

import logging
from io import BytesIO
from typing import Any

import pytest
import sqlalchemy as sa
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.engine import Engine

from tests import test_estimates_assumptions as assumption_tests
from tests.test_estimates_assumptions import (
    HAPPY,
    _accept_two,
    _register,
    by_title,
    events,
    proposed,
    redraft,
)
from tests.test_estimates_export_render import _sheet_rows
from tests.test_health import assert_problem
from tests.test_intake_extraction import BASE, PROPOSE, FakeGateway, drain, owner
from tests.test_opportunities import PSE, _add, _user

storage_dir = assumption_tests.storage_dir
_hidden_jobs = assumption_tests._hidden_jobs
gateway = assumption_tests.gateway
client = assumption_tests.client
_chat_profile = assumption_tests._chat_profile
_draft_profile = assumption_tests._draft_profile
_proposal_profile = assumption_tests._proposal_profile

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXPORTED = "estimates.estimate_version.exported"
SECTION_LABELS = {
    "functional": "Functional",
    "integration": "Integration",
    "data": "Data",
    "security": "Security",
    "non_functional": "Non-functional",
    "commercial": "Commercial",
}


def _export(client: TestClient, headers: dict[str, str], opp_id: str, fmt: str | None) -> Any:
    params = {} if fmt is None else {"format": fmt}
    return client.get(f"{BASE}/{opp_id}/estimate/export", headers=headers, params=params)


def carried_v2(
    client: TestClient, engine: Engine, db_url: str, gateway: FakeGateway
) -> tuple[dict[str, str], dict[str, Any]]:
    """Draft v2 of the SIX lines with 4 Assumptions: WMS (Condition) and Peak (8 h on
    "24/7 operations") carried from v1, SSO (Condition) and Test data owner (12 h,
    unallocated) not accepted; 4 Gaps with questions, 2 open and 2 converted, of which the
    open "SSO provider" Gap's question is superseded (so it is not exported)."""
    headers, opp = proposed(client, engine, db_url, gateway)
    _accept_two(client, headers, opp["id"])
    redraft(client, headers, opp["id"], db_url, gateway)
    gateway.replies = [by_title(HAPPY[2], HAPPY[3])]  # type: ignore[list-item]
    assert drain(db_url, PROPOSE) == ["succeeded"]
    with engine.begin() as conn:
        conn.execute(
            sa.text(
                "UPDATE gaps_clarification_questions SET status = 'superseded' WHERE gap_id = "
                "(SELECT id FROM gaps_gaps WHERE opportunity_id = :o AND title = 'SSO provider')"
            ),
            {"o": opp["id"]},
        )
    return headers, opp


def _expected_estimate(version: dict[str, Any]) -> list[tuple[Any, ...]]:
    """The Estimate sheet's rows below the column row, from the read model."""
    rows: list[tuple[Any, ...]] = []
    for section in version["sections"]:
        label = SECTION_LABELS[section["section"]]
        rows.append((label, None, None, None, None, None, None, None, None))
        for line in section["lines"]:
            mix = line["role_mix"]
            rows.append(
                (
                    label,
                    line["title"],
                    line["basis"],
                    mix["engineer"],
                    mix["project_manager"],
                    mix["qa"],
                    line["effort_hours"],
                    line["contingency_hours"],
                    line["total_hours"],
                )
            )
        sub = section["subtotal"]
        rows.append(
            (
                f"{label} subtotal",
                None,
                None,
                None,
                None,
                None,
                sub["effort_hours"],
                sub["contingency_hours"],
                sub["total_hours"],
            )
        )
    unallocated = version["unallocated_contingency_hours"]
    if unallocated > 0:
        rows.append(
            (
                "Unallocated contingency",
                None,
                None,
                None,
                None,
                None,
                None,
                unallocated,
                unallocated,
            )
        )
    t = version["totals"]
    rows.append(
        (
            "Total",
            None,
            None,
            None,
            None,
            None,
            t["effort_hours"],
            t["contingency_hours"],
            t["total_hours"],
        )
    )
    return rows


def test_the_xlsx_export_matches_the_estimate_register_and_questions(
    client: TestClient,
    sync_engine: Engine,
    db_url: str,
    gateway: FakeGateway,
    caplog: pytest.LogCaptureFixture,
) -> None:
    headers, opp = carried_v2(client, sync_engine, db_url, gateway)
    version = _register(client, headers, opp["id"])
    shown = [
        {"id": g["question"]["id"], "row_version": g["question"]["row_version"]}
        for g in client.get(f"{BASE}/{opp['id']}/gaps", headers=headers).json()["items"]
        if g["status"] == "open" and g["question"] and g["question"]["status"] == "drafted"
    ]
    approved = client.post(
        f"{BASE}/{opp['id']}/clarification-questions/approve-all", headers=headers, json=shown
    )
    assert approved.json() == {"count": 1}, approved.text  # the open "Test data owner" Gap's
    with sync_engine.begin() as conn:  # and a converted Gap's, approved before it converted
        conn.execute(
            sa.text(
                "UPDATE gaps_clarification_questions SET status = 'approved', "
                "approved_by = gen_random_uuid(), approved_at = now() WHERE gap_id = "
                "(SELECT id FROM gaps_gaps WHERE opportunity_id = :o AND title = :t)"
            ),
            {"o": opp["id"], "t": "WMS version unknown"},
        )
    gaps = client.get(f"{BASE}/{opp['id']}/gaps", headers=headers).json()["items"]

    with caplog.at_level(logging.DEBUG):
        resp = _export(client, headers, opp["id"], "xlsx")

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == XLSX
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=")
    assert disposition.endswith('-estimate-v2.xlsx"')
    assert disposition.isascii()
    workbook = load_workbook(BytesIO(resp.content))
    assert workbook.sheetnames == ["Estimate", "Assumptions", "Clarification Questions"]
    for sheet in workbook.worksheets:
        head = [r[0] for r in _sheet_rows(sheet)[:3]]
        assert head[0] == opp["title"]
        assert head[1] == "Estimate v2 · Draft — not submitted"
        assert head[2] is not None and head[2].startswith("Exported ")
        assert head[2].endswith(" UTC")

    estimate = _sheet_rows(workbook["Estimate"])[5:]
    assert estimate == _expected_estimate(version)
    assert version["unallocated_contingency_hours"] == 12.0
    assert estimate[-1][6:] == (108.5, 20.0, 128.5)

    register = _sheet_rows(workbook["Assumptions"])[5:]
    groups = version["assumptions"]
    items = [r for r in register if r[0] in ("Condition", "Contingency")]
    assert len(items) == 4
    by_gap = {r[4]: r for r in items}
    for a in groups["conditions"] + groups["contingencies"]:
        row = by_gap[a["origin"]["title"]]
        assert row[1] == a["wording"]
        assert row[2] == a["amount_hours"]
        if a["accepted_by"] is None:
            assert row[5] == "Not accepted"
        else:
            assert row[5].startswith(f"Accepted by {a['accepted_by']['name']}, ")
        carried = a["carried_from_version"]
        assert row[6] == (None if carried is None else f"Carried from v{carried}")
    assert by_gap["Peak order volume"][3] == "24/7 operations"
    assert by_gap["Test data owner"][3] == "Unallocated contingency"
    assert by_gap["WMS version unknown"][6] == "Carried from v1"
    assert by_gap["SSO provider"][5] == "Not accepted"

    questions = _sheet_rows(workbook["Clarification Questions"])[5:]
    expected = [
        (
            g["title"],
            g["impact"].capitalize(),
            g["question"]["topic"],
            g["question"]["text"],
            {"drafted": "Draft", "approved": "Approved"}[g["question"]["status"]],
            g["status"].capitalize(),
        )
        for g in gaps
        if g["question"] is not None
        and g["question"]["status"] in ("drafted", "approved")
        and g["status"] in ("open", "converted")
    ]
    assert len(gaps) == 4
    assert len(expected) == 3  # the superseded question is left out
    assert "SSO provider" not in {q[0] for q in questions}
    assert questions == expected
    assert sorted(q[4] for q in questions) == ["Approved", "Approved", "Draft"]
    assert {q[5] for q in questions} == {"Open", "Converted"}

    (event,) = events(sync_engine, opp["id"], EXPORTED)
    assert (event["subject_type"], str(event["subject_id"])) == (
        "estimates.estimate_version",
        version["id"],
    )
    assert event["actor_type"] == "user"
    assert event["payload"] == {"version_id": version["id"], "version": 2, "format": "xlsx"}

    # Privacy: no line, Assumption, Gap or question text in logs or the trace.
    logged = caplog.text + str([r.__dict__ for r in caplog.records])
    trace = str(event)
    for leak in ("estimate assumes", "Contingency for", "WMS version", "24/7", "Could you tell"):
        assert leak not in logged
        assert leak not in trace


def test_the_docx_export_has_the_same_content(
    client: TestClient, sync_engine: Engine, db_url: str, gateway: FakeGateway
) -> None:
    headers, opp = carried_v2(client, sync_engine, db_url, gateway)
    version = _register(client, headers, opp["id"])

    resp = _export(client, headers, opp["id"], "docx")

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == DOCX
    assert resp.headers["content-disposition"].endswith('-estimate-v2.docx"')
    doc = Document(BytesIO(resp.content))
    header = [p.text for p in doc.sections[0].header.paragraphs]
    assert header[:2] == [opp["title"], "Estimate v2 · Draft — not submitted"]
    assert header[2].startswith("Exported ")
    estimate, register, questions = (
        [tuple(c.text for c in row.cells) for row in table.rows] for table in doc.tables
    )

    def text(value: Any) -> str:
        return "" if value is None else f"{value:.1f}" if isinstance(value, float) else str(value)

    assert estimate[1:] == [tuple(text(v) for v in r) for r in _expected_estimate(version)]
    assert len([r for r in register if r[0] in ("Condition", "Contingency")]) == 4
    assert len(questions) == 4  # the column row and 3 questions (one is superseded)
    (event,) = events(sync_engine, opp["id"], EXPORTED)
    assert event["payload"]["format"] == "docx"


def test_no_version_means_nothing_to_export(client: TestClient, sync_engine: Engine) -> None:
    headers, opp = owner(client, sync_engine)

    resp = _export(client, headers, opp["id"], "xlsx")

    assert_problem(resp.json(), 409, "estimate_not_found")
    assert events(sync_engine, opp["id"], EXPORTED) == []


def test_who_may_export(
    client: TestClient, sync_engine: Engine, db_url: str, gateway: FakeGateway
) -> None:
    headers, opp = proposed(client, sync_engine, db_url, gateway)
    rep, rep_id = _user(client, sync_engine, "Rep Person", "sales_representative")
    current = client.get(f"{BASE}/{opp['id']}", headers=headers).json()
    assert _add(client, headers, current, rep_id).status_code == 200
    reader, _ = _user(client, sync_engine, "Head Person", "head_of_delivery")
    outsider, _ = _user(client, sync_engine, "Outsider", PSE)
    colleague, colleague_id = _user(client, sync_engine, "Colleague Person", PSE)
    current = client.get(f"{BASE}/{opp['id']}", headers=headers).json()
    assert _add(client, headers, current, colleague_id).status_code == 200

    assert client.get(f"{BASE}/{opp['id']}/estimate", headers=rep).json()["can_export"] is False
    assert _register(client, headers, opp["id"]) is not None
    assert client.get(f"{BASE}/{opp['id']}/estimate", headers=headers).json()["can_export"]
    assert_problem(_export(client, rep, opp["id"], "xlsx").json(), 403, "forbidden")
    assert_problem(_export(client, reader, opp["id"], "xlsx").json(), 403, "forbidden")
    assert_problem(_export(client, outsider, opp["id"], "xlsx").json(), 404, "not_found")
    assert events(sync_engine, opp["id"], EXPORTED) == []
    assert client.get(f"{BASE}/{opp['id']}/estimate", headers=colleague).json()["can_export"]
    assert _export(client, colleague, opp["id"], "xlsx").status_code == 200
    assert _export(client, headers, opp["id"], "xlsx").status_code == 200
    assert len(events(sync_engine, opp["id"], EXPORTED)) == 2


@pytest.mark.parametrize("fmt", ["pdf", "", None])
def test_a_bad_format_is_refused(client: TestClient, sync_engine: Engine, fmt: str | None) -> None:
    headers, opp = owner(client, sync_engine)

    resp = _export(client, headers, opp["id"], fmt)

    assert_problem(resp.json(), 422, "validation_error")
