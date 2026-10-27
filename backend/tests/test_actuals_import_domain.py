"""Actuals import rules (Story 10.2), pure: header detection, blank rows, reading one row into
an Actual's fields with "Row N: ..." problems, cost in major units, the repeated-file message,
the xlsx/csv content checks, the template, and the sandbox child on real files."""

import asyncio
import io
import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.actuals.adapters import import_cli
from app.modules.actuals.adapters.template_xlsx import (
    TemplateDocument,
    TemplateLine,
    render_template,
)
from app.modules.actuals.domain.actuals import CAUSE_REQUIRED, ActualKind, VarianceCause
from app.modules.actuals.domain.imports import (
    COLUMNS,
    ImportErrorCode,
    ImportReadError,
    cell_text,
    currency_exponent,
    duplicate_message,
    is_blank,
    read_row,
    rule_text,
    table_rows,
)
from app.platform.parsing import runner
from app.platform.parsing.rules import ParseError
from app.platform.upload_validation import content_matches

LINE = uuid4()
A1 = uuid4()
HEADERS = [c.header for c in COLUMNS]


def _read(**values: str) -> tuple[object, list[str]]:
    return read_row(7, values, baseline_version=3, line_ids={LINE}, assumptions={1: A1})


def test_cell_text() -> None:
    assert [cell_text(v) for v in (None, 41, 41.0, 41.5, True, "  x ")] == [
        "",
        "41",
        "41",
        "41.5",
        "TRUE",
        "x",
    ]


def test_header_row_is_found_and_blank_rows_dropped() -> None:
    table = [
        (1, ["Actuals for v3"]),
        (2, ["line id", "  ACTUAL   H ", "Note"]),
        (3, [str(LINE), 36, None]),
        (4, [None, None]),
        (5, ["", "", "late"]),
    ]
    assert table_rows(table) == [
        {"row": 3, "values": {"line_id": str(LINE), "actual_hours": "36"}},
        {"row": 5, "values": {"note": "late"}},
    ]


def test_missing_columns() -> None:
    with pytest.raises(ImportReadError) as raised:
        table_rows([(1, ["Name", "Hours"]), (2, ["L1", 3])])
    assert raised.value.code == ImportErrorCode.MISSING_COLUMNS


def test_blank_means_no_actual_field() -> None:
    assert is_blank({"line_id": str(LINE), "line": "L1", "estimate_hours": "36"})
    assert not is_blank({"line_id": str(LINE), "note": "x"})


def test_a_full_row() -> None:
    entry, problems = _read(
        line_id=str(LINE),
        kind="Milestone",
        milestone_label="Sprint 1",
        actual_hours="12.5",
        duration_days="3",
        cost="1,250.50",
        currency="eur",
        variance_cause="Estimation error",
        failed_assumptions="a1, A1",
        note="n",
        correction_reason="r",
    )
    assert problems == []
    assert entry is not None
    assert entry.kind == ActualKind.MILESTONE  # type: ignore[attr-defined]
    assert entry.effort_hours == Decimal("12.5")  # type: ignore[attr-defined]
    assert entry.cost_minor == 125050  # type: ignore[attr-defined]
    assert entry.currency == "EUR"  # type: ignore[attr-defined]
    assert entry.variance_cause == VarianceCause.ESTIMATION_ERROR  # type: ignore[attr-defined]
    assert entry.failed_assumption_ids == [A1]  # type: ignore[attr-defined]


def test_kind_defaults_to_final_and_cause_takes_its_value() -> None:
    entry, _ = _read(line_id=str(LINE), actual_hours="41", variance_cause="scope_change")
    assert entry is not None
    assert entry.kind == ActualKind.FINAL  # type: ignore[attr-defined]
    assert entry.variance_cause == VarianceCause.SCOPE_CHANGE  # type: ignore[attr-defined]


@pytest.mark.parametrize(
    ("values", "problem"),
    [
        ({"actual_hours": "5"}, "Row 7: line ID missing"),
        ({"line_id": str(uuid4()), "actual_hours": "5"}, "Row 7: line ID not in Baseline v3"),
        ({"line_id": "L1", "actual_hours": "5"}, "Row 7: line ID not in Baseline v3"),
        ({"line_id": str(LINE), "note": "x"}, "Row 7: actual hours missing"),
        ({"line_id": str(LINE), "actual_hours": "x"}, "Row 7: actual hours must be a number"),
        (
            {"line_id": str(LINE), "actual_hours": "5", "kind": "partial"},
            "Row 7: kind must be Final or Milestone",
        ),
        (
            {"line_id": str(LINE), "actual_hours": "5", "failed_assumptions": "A2"},
            "Row 7: A2 isn't in the Baseline v3 Assumptions Register",
        ),
        (
            {"line_id": str(LINE), "actual_hours": "5", "cost": "1.5", "currency": "JPY"},
            "Row 7: cost has more decimal places than JPY allows",
        ),
    ],
)
def test_row_problems(values: dict[str, str], problem: str) -> None:
    entry, problems = _read(**values)
    assert entry is None
    assert problem in problems


def test_currency_exponents() -> None:
    assert [currency_exponent(c) for c in ("EUR", "JPY", "BHD", "CLF", None)] == [2, 0, 3, 4, 2]


def test_rule_text() -> None:
    assert rule_text("variance_cause", CAUSE_REQUIRED) == (
        "variance cause missing for a final Actual"
    )
    assert rule_text("reason", "anything") == "correction reason missing"
    assert rule_text("currency", "Currency is an ISO-4217 code.") == "currency is an ISO-4217 code"


def test_duplicate_message() -> None:
    assert duplicate_message(datetime(2026, 10, 14, 9, tzinfo=UTC)) == (
        "This file was already imported on 14 Oct"
    )


def test_xlsx_and_csv_content_checks() -> None:
    book = io.BytesIO()
    Workbook().save(book)
    assert content_matches(".xlsx", io.BytesIO(book.getvalue()))
    docx = io.BytesIO()
    with zipfile.ZipFile(docx, "w") as archive:
        archive.writestr("word/document.xml", "x")
    assert not content_matches(".xlsx", io.BytesIO(docx.getvalue()))
    assert not content_matches(".xlsx", io.BytesIO(b"%PDF-1.7"))
    assert content_matches(".csv", io.BytesIO("Line ID,Actual h\nä,1\n".encode()))
    assert not content_matches(".csv", io.BytesIO(b"\x00\xff"))


def test_template_text_is_never_a_formula() -> None:
    data = render_template(
        TemplateDocument(
            estimate_version=3,
            lines=[TemplateLine(str(LINE), "L1", "=HYPERLINK(1)", "", 36.0)],
            kinds=["Final", "Milestone"],
            causes=["Other"],
            assumptions=[("A1", "Condition", "Volumes as stated")],
        )
    )
    book = load_workbook(io.BytesIO(data))
    sheet = book["Actuals"]
    assert [c.value for c in sheet[1]] == HEADERS
    assert sheet["C2"].value == "=HYPERLINK(1)"
    assert sheet["C2"].data_type == "s"
    assert book["Lists"]["A1"].value == "Estimate v3"


def _run(path: Path, ext: str) -> runner.ChildOutput:
    return asyncio.run(
        runner.run_parse(path, ext, timeout_s=60, max_memory_mb=512, child_module=imports_child())
    )


def imports_child() -> str:
    return import_cli.__name__


def test_the_child_reads_xlsx_and_csv(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(HEADERS)
    sheet.append([str(LINE), "L1", None, None, 36, None, None, 41.5])
    xlsx = tmp_path / "blob-without-extension"  # as stored, by SHA-256
    book.save(xlsx)
    output = _run(xlsx, ".xlsx")
    assert output.parser == import_cli.PARSER_XLSX
    assert '"41.5"' not in output.text and "41.5" in output.text

    csv_file = tmp_path / "a.csv"
    csv_file.write_bytes("﻿Line ID;Actual h\nx;3\n".encode())
    output = _run(csv_file, ".csv")
    assert output.parser == import_cli.PARSER_CSV
    assert output.text == '[[1, ["Line ID", "Actual h"]], [2, ["x", "3"]]]'


def test_the_child_fails_on_a_corrupt_workbook(tmp_path: Path) -> None:
    bad = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(bad, "w") as archive:
        archive.writestr("xl/workbook.xml", "not a workbook")
    with pytest.raises(ParseError) as raised:
        _run(bad, ".xlsx")
    assert raised.value.code.value == "unreadable"


@pytest.mark.parametrize(
    ("typed", "hours"),
    [
        ("41,5", "41.5"),
        ("1234,56", "1234.56"),
        ("1,250.50", "1250.50"),
        ("1,234", "1234"),
        ("12,345,678", "12345678"),
        ("41.5", "41.5"),
    ],
)
def test_numbers_with_commas(typed: str, hours: str) -> None:
    entry, problems = _read(line_id=str(LINE), actual_hours=typed, variance_cause="other")
    assert problems == []
    assert entry is not None
    assert entry.effort_hours == Decimal(hours)  # type: ignore[attr-defined]


@pytest.mark.parametrize("typed", ["1,2,3", "12,34.5", "1,23", "4,1.5", ","])
def test_badly_grouped_commas_are_not_numbers(typed: str) -> None:
    if typed == "1,23":  # a single comma, no point: a decimal comma
        entry, _ = _read(line_id=str(LINE), actual_hours=typed)
        assert entry is not None
        assert entry.effort_hours == Decimal("1.23")  # type: ignore[attr-defined]
        return
    _, problems = _read(line_id=str(LINE), actual_hours=typed)
    assert problems == ["Row 7: actual hours must be a number"]


def test_a_formula_float_reads_as_typed() -> None:
    assert cell_text(41.300000000000004) == "41.3"
    assert cell_text(0.1 + 0.2) == "0.3"
    assert cell_text(1234.5678) == "1234.5678"


def test_too_many_failed_assumptions() -> None:
    many = {n: uuid4() for n in range(1, 53)}
    labels = ", ".join(f"A{n}" for n in many)
    entry, problems = read_row(
        7,
        {"line_id": str(LINE), "actual_hours": "5", "failed_assumptions": labels},
        baseline_version=3,
        line_ids={LINE},
        assumptions=many,
    )
    assert entry is None
    assert problems == ["Row 7: at most 50 failed Assumptions per row"]


def test_a_too_long_own_reason_is_a_row_problem() -> None:
    entry, problems = _read(line_id=str(LINE), actual_hours="5", correction_reason="x" * 1001)
    assert entry is None
    assert problems == ["Row 7: correction reason is longer than 1,000 characters"]


def test_the_child_reads_quoted_line_breaks(tmp_path: Path) -> None:
    csv_file = tmp_path / "blob"
    csv_file.write_bytes(b'Line ID,Actual h,Note\r\nx,3,"two\r\nlines"\r\ny,4,one\r\n')
    output = _run(csv_file, ".csv")
    assert output.text == (
        '[[1, ["Line ID", "Actual h", "Note"]], [2, ["x", "3", "two\\r\\nlines"]], '
        '[4, ["y", "4", "one"]]]'
    )


def test_the_child_finds_the_delimiter_below_a_title(tmp_path: Path) -> None:
    csv_file = tmp_path / "blob"
    csv_file.write_bytes(b"Actuals for v3, October\nLine ID;Actual h;Note\nx;41,5;n\n")
    output = _run(csv_file, ".csv")
    assert output.text == (
        '[[1, ["Actuals for v3, October"]], [2, ["Line ID", "Actual h", "Note"]], '
        '[3, ["x", "41,5", "n"]]]'
    )
