"""The Actuals import child process (Story 10.2):
`python -m app.modules.actuals.adapters.import_cli <blob path> <ext>`.

Runs in the platform's parse sandbox (`app.platform.parsing.runner`: wall-clock, CPU and
memory limits, an allowlisted environment), never in the api process. First thing it silences
warnings and logging and sets its limits (the platform child's own helpers); only then is the
reader imported. It reads the first worksheet of an `.xlsx` (openpyxl, `read_only` and
`data_only`: cached values, never formulas) or a UTF-8 `.csv` (stdlib) and writes, as JSON on
stdout, `[[row number, [cell, ...]], ...]` with at most `MAX_TABLE_ROWS` rows and
`MAX_COLUMNS` cells per row (cells are strings, numbers, booleans or null; dates as ISO
text). The last stderr line is `parser=<name>` on success, or the failure code
(`unreadable`, `too_large_output`) with exit `EXIT_PARSE_ERROR`. No content, file names or
tracebacks are written to stderr.
"""

import math
import sys

from app.platform.parsing.cli import EXIT_PARSE_ERROR, EXIT_USAGE, _apply_limits, _fail, _quiet

MAX_TABLE_ROWS = 2100
"""Header search rows plus `domain.imports.MAX_ROWS`, with room: past it, `too_large_output`."""
MAX_COLUMNS = 40
PARSER_XLSX = "actuals-xlsx@openpyxl-3.1.5"
PARSER_CSV = "actuals-csv@stdlib"
UNREADABLE = "unreadable"
TOO_LARGE = "too_large_output"

__all__ = ["EXIT_PARSE_ERROR", "main"]


class _TooLargeError(Exception):
    pass


def _cell(value: object) -> object:
    if isinstance(value, float) and not math.isfinite(value):  # JSON has no NaN or infinity
        return str(value)
    if value is None or isinstance(value, bool | int | float | str):
        return value
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):  # date, datetime, time
        return str(isoformat())
    return str(value)


def _trimmed(cells: list[object]) -> list[object]:
    while cells and (cells[-1] is None or cells[-1] == ""):
        cells.pop()
    return cells


def _xlsx(path: str) -> list[list[object]]:
    from io import BytesIO
    from pathlib import Path

    from openpyxl import load_workbook

    # A stored blob has no extension, and openpyxl refuses a path without `.xlsx`: read
    # the bytes and hand over a file object instead.
    data = BytesIO(Path(path).read_bytes())
    workbook = load_workbook(data, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        table: list[list[object]] = []
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = _trimmed([_cell(v) for v in values[:MAX_COLUMNS]])
            if not cells:
                continue
            if len(table) >= MAX_TABLE_ROWS:
                raise _TooLargeError
            table.append([number, cells])
        return table
    finally:
        workbook.close()


def _csv(path: str) -> list[list[object]]:
    import csv
    import io
    from pathlib import Path

    from app.modules.actuals.domain.imports import HEADER_SEARCH_ROWS

    text = Path(path).read_bytes().decode("utf-8-sig")
    # The delimiter from the first lines, so a title row above the header doesn't hide it.
    head = "\n".join(text.splitlines()[:HEADER_SEARCH_ROWS])
    delimiter = max((",", ";", "\t"), key=head.count) if head else ","
    table: list[list[object]] = []
    # newline="": a quoted cell may hold line breaks; each row is numbered by its first line.
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    number = 1
    for values in reader:
        cells = _trimmed([v.strip() for v in values[:MAX_COLUMNS]])
        start, number = number, reader.line_num + 1
        if not cells:
            continue
        if len(table) >= MAX_TABLE_ROWS:
            raise _TooLargeError
        table.append([start, cells])
    return table


def main(argv: list[str]) -> int:
    _quiet()
    _apply_limits()
    if len(argv) != 2:
        sys.stderr.write("usage")
        return EXIT_USAGE
    path, ext = argv
    try:
        import json

        if ext == ".xlsx":
            table, parser = _xlsx(path), PARSER_XLSX
        elif ext == ".csv":
            table, parser = _csv(path), PARSER_CSV
        else:
            return _fail(UNREADABLE)
        encoded = json.dumps(table, ensure_ascii=False, allow_nan=False).encode("utf-8")
    except _TooLargeError:
        return _fail(TOO_LARGE)
    except Exception:  # corrupt, encrypted, not UTF-8, MemoryError under the limit
        return _fail(UNREADABLE)
    sys.stdout.buffer.write(encoded)
    sys.stdout.buffer.flush()
    sys.stderr.write(f"\nparser={parser}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
