"""The `.xlsx` Actuals import template (Story 10.2). Pure: a `TemplateDocument` in, workbook
bytes out.

Sheet "Actuals" has the header row (`domain.imports.COLUMNS`) and one row per Baseline line:
its Line ID (the Estimate line's UUID), `L{n}`, Integration Type, Work Package and Estimate h,
then empty cells for the Actual. Sheet "Lists" names the accepted kinds and variance causes
and the Baseline version's Assumptions as `A{n}`. Every text cell is stored as a string, never
as a formula."""

from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.actuals.domain.imports import COLUMNS

HOURS_FORMAT = "0.0"
_XML_ILLEGAL = str.maketrans({i: None for i in range(32) if i not in (9, 10, 13)})
_WIDTHS = {"line_id": 38, "integration_type": 28, "work_package": 24, "note": 40}


@dataclass(frozen=True, slots=True)
class TemplateLine:
    line_id: str
    label: str
    integration_type: str
    work_package: str
    estimate_hours: float


@dataclass(frozen=True, slots=True)
class TemplateDocument:
    estimate_version: int
    lines: Sequence[TemplateLine]
    kinds: Sequence[str]
    causes: Sequence[str]
    assumptions: Sequence[tuple[str, str, str]]
    """`(A{n}, kind, wording)` in register order."""


def _clean(text: str) -> str:
    return text.translate(_XML_ILLEGAL)


def _append(sheet: Worksheet, values: Sequence[Any]) -> tuple[Any, ...]:
    sheet.append([_clean(v) if isinstance(v, str) else v for v in values])
    cells: tuple[Any, ...] = tuple(sheet[sheet.max_row])
    for cell in cells:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # text, never a formula
    return cells


def render_template(document: TemplateDocument) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.title = "Actuals"
    for cell in _append(sheet, [c.header for c in COLUMNS]):
        cell.font = Font(bold=True)
    for line in document.lines:
        cells = _append(
            sheet,
            [
                line.line_id,
                line.label,
                line.integration_type or None,
                line.work_package or None,
                line.estimate_hours,
                *([None] * (len(COLUMNS) - 5)),
            ],
        )
        cells[4].number_format = HOURS_FORMAT
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = _WIDTHS.get(
            column.key, max(12, len(column.header) + 2)
        )
    sheet.freeze_panes = "C2"

    lists = workbook.create_sheet("Lists")
    for cell in _append(lists, [f"Estimate v{document.estimate_version}"]):
        cell.font = Font(bold=True)
    lists.append([])
    for cell in _append(lists, ["Kind"]):
        cell.font = Font(bold=True)
    for kind in document.kinds:
        _append(lists, [kind])
    lists.append([])
    for cell in _append(lists, ["Variance cause"]):
        cell.font = Font(bold=True)
    for cause in document.causes:
        _append(lists, [cause])
    lists.append([])
    for cell in _append(lists, ["Failed Assumptions", "Kind", "Wording"]):
        cell.font = Font(bold=True)
    for label, kind, wording in document.assumptions:
        _append(lists, [label, kind, wording])
    lists.column_dimensions["A"].width = 24
    lists.column_dimensions["C"].width = 80

    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def file_name(estimate_version: int) -> str:
    return f"actuals-template-estimate-v{estimate_version}.xlsx"
