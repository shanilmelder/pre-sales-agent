"""The `.docx` renderer of an Estimate export (Story 8.8). Pure: an `ExportDocument` in,
document bytes out. The export header is the page header (so it is on every page); each
table is a heading and a Word table, ready to paste into the proposal."""

from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import _Cell

from app.modules.estimates.application.export_content import (
    Cell,
    ExportDocument,
    Table,
    clean_text,
)

BOLD_KINDS = frozenset({"group", "subtotal", "total"})
TABLE_STYLE = "Table Grid"


def cell_text(value: Cell) -> str:
    """A cell as text: hours with one decimal place, whole numbers as they are."""
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.1f}"
    return clean_text(str(value))


def _fill(cell: _Cell, text: str, *, bold: bool, right: bool) -> None:
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(clean_text(text))
    run.bold = bold
    run.font.size = Pt(9)
    if right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _table(document: DocumentObject, table: Table) -> None:
    document.add_heading(clean_text(table.title), level=1)
    grid = document.add_table(rows=1, cols=len(table.columns))
    grid.style = TABLE_STYLE
    # The column row repeats at the top of every page the table runs onto.
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    grid.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, name in enumerate(table.columns):
        _fill(grid.rows[0].cells[index], name, bold=True, right=index in table.number_columns)
    for row in table.rows:
        cells = grid.add_row().cells
        for index, value in enumerate(row.cells):
            _fill(
                cells[index],
                cell_text(value),
                bold=row.kind in BOLD_KINDS,
                right=index in table.number_columns,
            )


def render_docx(document: ExportDocument) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    width, height = section.page_width, section.page_height
    if width is not None and height is not None and width < height:
        section.page_width, section.page_height = height, width
    header = section.header
    first, *rest = document.header
    header.paragraphs[0].text = clean_text(first)
    header.paragraphs[0].runs[0].bold = True
    for text in rest:
        header.add_paragraph(clean_text(text))
    for table in document.tables:
        _table(doc, table)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
