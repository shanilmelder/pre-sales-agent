"""The `.xlsx` renderer of an Estimate export (Story 8.8). Pure: an `ExportDocument` in,
workbook bytes out. One sheet per table, each starting with the export header."""

from collections.abc import Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.estimates.application.export_content import (
    Cell,
    ExportDocument,
    Table,
    clean_text,
)

HOURS_FORMAT = "0.0"
WIDE = 48
NARROW = 14
BOLD_KINDS = frozenset({"group", "subtotal", "total"})


def _sheet_title(title: str) -> str:
    return title[:31]  # Excel's limit


def _append(sheet: Worksheet, values: Sequence[Cell]) -> tuple[Any, ...]:
    """Append a row; text is cleaned and always stored as a string, never as a formula (a
    title starting with `=` stays text)."""
    sheet.append([clean_text(v) if isinstance(v, str) else v for v in values])
    cells: tuple[Any, ...] = tuple(sheet[sheet.max_row])
    for cell in cells:
        if isinstance(cell.value, str):
            cell.data_type = "s"
    return cells


def _write(sheet: Worksheet, header: tuple[str, ...], table: Table) -> None:
    for text in header:
        _append(sheet, [text])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    for cell in _append(sheet, table.columns):
        cell.font = Font(bold=True)
    head_row = sheet.max_row
    for row in table.rows:
        current = _append(sheet, row.cells)
        for index, cell in enumerate(current):
            if index in table.number_columns:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(cell.value, float):
                    cell.number_format = HOURS_FORMAT
            else:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row.kind in BOLD_KINDS:
                cell.font = Font(bold=True)
    for index, name in enumerate(table.columns):
        width = NARROW if index in table.number_columns else WIDE
        if index == 0:
            width = 28
        sheet.column_dimensions[get_column_letter(index + 1)].width = max(width, len(name) + 2)
    sheet.freeze_panes = f"A{head_row + 1}"


def render_xlsx(document: ExportDocument) -> bytes:
    workbook = Workbook()
    first = workbook.active
    assert isinstance(first, Worksheet)
    for index, table in enumerate(document.tables):
        sheet = first if index == 0 else workbook.create_sheet()
        sheet.title = _sheet_title(table.title)
        _write(sheet, document.header, table)
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()
