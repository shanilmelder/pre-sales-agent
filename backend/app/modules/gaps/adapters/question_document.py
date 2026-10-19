"""The customer-facing document of approved Clarification Questions (Story 4.6). Pure: no DB.

`build_document` groups the questions by topic (topics A to Z, ignoring case; within a topic,
by Gap rank) and numbers them 1 to N across the document. It carries only what the customer
may see: the customer's name, a date, the topic headings and the question wording; never a
Gap's title, reason, impact, trigger, Knowledge Source or Evidence. `plain_text`,
`render_docx` and `render_eml` lay the same content out as text to paste, a Word document and
an unsent email draft (`X-Unsent: 1`, no recipient). Nothing here sends anything.
"""

import re
import unicodedata
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO

from docx import Document
from docx.shared import Pt

MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
SLUG_MAX = 60
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EML_MEDIA_TYPE = "message/rfc822"
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass(frozen=True, slots=True)
class ExportQuestion:
    """One approved question as the export needs it. `rank` orders it within its topic (the
    Gap's rank: impact, then age)."""

    topic: str = field(repr=False)
    text: str = field(repr=False)
    rank: tuple[int, str, str]


@dataclass(frozen=True, slots=True)
class NumberedQuestion:
    number: int
    text: str = field(repr=False)


@dataclass(frozen=True, slots=True)
class TopicGroup:
    topic: str = field(repr=False)
    questions: tuple[NumberedQuestion, ...]


@dataclass(frozen=True, slots=True)
class QuestionDocument:
    customer_name: str = field(repr=False)
    on: date
    groups: tuple[TopicGroup, ...]

    @property
    def title(self) -> str:
        return f"Clarification questions for {self.customer_name}"

    @property
    def date_label(self) -> str:
        return day_label(self.on)


def day_label(on: date) -> str:
    """'14 Oct 2026'."""
    return f"{on.day} {MONTHS[on.month - 1]} {on.year}"


def clean_text(text: str) -> str:
    """`text` without the control characters XML can't hold (tab, newline and carriage return
    stay)."""
    return _XML_ILLEGAL.sub("", text)


def build_document(
    *, customer_name: str, on: date, questions: Sequence[ExportQuestion]
) -> QuestionDocument:
    """The questions grouped by topic (A to Z, ignoring case and surrounding spacing), each
    group ordered by rank, numbered 1 to N in that order."""
    by_topic: dict[str, list[ExportQuestion]] = {}
    names: dict[str, str] = {}
    for question in questions:
        key = " ".join(question.topic.casefold().split())
        by_topic.setdefault(key, []).append(question)
        names.setdefault(key, " ".join(question.topic.split()))
    groups: list[TopicGroup] = []
    number = 0
    for key in sorted(by_topic):
        numbered: list[NumberedQuestion] = []
        for question in sorted(by_topic[key], key=lambda q: q.rank):
            number += 1
            numbered.append(NumberedQuestion(number=number, text=question.text.strip()))
        groups.append(TopicGroup(topic=names[key], questions=tuple(numbered)))
    return QuestionDocument(customer_name=customer_name.strip(), on=on, groups=tuple(groups))


def plain_text(document: QuestionDocument) -> str:
    """The document as plain text: the title and date, then each topic and its numbered
    questions, separated by blank lines."""
    lines = [clean_text(document.title), document.date_label]
    for group in document.groups:
        lines.append("")
        lines.append(clean_text(group.topic))
        for question in group.questions:
            lines.append(f"{question.number}. {clean_text(question.text)}")
    return "\n".join(lines) + "\n"


def render_docx(document: QuestionDocument) -> bytes:
    """The document as a Word file: a title, the date, a heading per topic and its numbered
    questions."""
    doc = Document()
    doc.add_heading(clean_text(document.title), level=0)
    doc.add_paragraph(document.date_label)
    for group in document.groups:
        doc.add_heading(clean_text(group.topic), level=1)
        for question in group.questions:
            paragraph = doc.add_paragraph()
            number = paragraph.add_run(f"{question.number}. ")
            number.bold = True
            run = paragraph.add_run(clean_text(question.text))
            run.font.size = Pt(11)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def email_subject(document: QuestionDocument) -> str:
    return " ".join(f"Clarification questions — {document.customer_name}".split())


def render_eml(document: QuestionDocument) -> bytes:
    """The document as an unsent plain-text email draft: a subject, no recipient, and
    `X-Unsent: 1` so mail clients open it as a draft to send."""
    message = EmailMessage(policy=SMTP)
    message["Subject"] = email_subject(document)
    message["X-Unsent"] = "1"
    message.set_content(plain_text(document), subtype="plain", charset="utf-8")
    return message.as_bytes()


def file_name(customer_name: str, on: date, extension: str) -> str:
    """`clarification-questions-<customer>-<yyyy-mm-dd>.<extension>`, ASCII only."""
    ascii_name = (
        unicodedata.normalize("NFKD", customer_name).encode("ascii", "ignore").decode("ascii")
    )
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.lower()).strip("-")[:SLUG_MAX].strip("-")
    stem = "clarification-questions" + (f"-{slug}" if slug else "")
    return f"{stem}-{on.isoformat()}.{extension}"
